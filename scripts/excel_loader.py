#!/usr/bin/env python3
"""
excel_loader.py

Mengambil file .xlsx dari folder sumber di Nextcloud (via WebDAV), membaca
sheet yang dikenal kontraknya, lalu menyimpannya LANGSUNG ke PostgreSQL —
menggantikan Airbyte. Satu image, dua dataset (LOADER_DATASET):

  finance (default, service excel-loader) — docs/adr/0005
    <Kantor>/Finance/*.xlsx, sheet REKAP   -> raw_finance_rekap_<kantor_id>
                             sheet RINCIAN -> raw_finance_rincian_<kantor_id>
    kantor_id = komponen path tepat di atas 'Finance': .../A1/Finance -> a1.

  capil (service capil-loader) — docs/adr/0006
    <Kantor>/data.xlsx (CAPIL_FILE), satu sheet     -> raw_<kantor_id>
    kantor_id = komponen path terakhir: .../A1 -> a1.
    Kolomnya tidak ditulis di kode: kontraknya dibaca dari tabel raw kantor itu
    (bentuk buatan Airbyte), atau dari tabel capil kantor lain untuk kantor baru.

Tabel tujuan sama persis dengan yang dulu ditulis Airbyte, supaya dbt, hist_*
dan Laporan Rekap Bulanan tidak perlu berubah.

Setiap muatan adalah satu "tarikan" gaya Airbyte mode Append: baris baru dengan
_airbyte_generation_id = MAX + 1, bukan menimpa. Kolom _airbyte_* sengaja
dipertahankan — filter "tarikan terakhir" di staging, macro hist_pulls, dan
report_summary.py semuanya bergantung pada kolom itu.

Mode:
  python excel_loader.py           -- jalankan sekali lalu keluar
  python excel_loader.py A2/Finance [B1/Finance ...]
                                   -- jalankan sekali untuk folder tertentu saja
                                      (capil: python excel_loader.py A2 B1)
  python excel_loader.py --watch   -- loop berkala (SCHEDULE_INTERVAL_MINUTES)
  python excel_loader.py --serve   -- server kontrol HTTP (job_control) di gisnet,
                                      dipakai bot Telegram untuk memicu run

Path sumber:
  NEXTCLOUD_SOURCE_PATHS (finance) / CAPIL_SOURCE_PATHS (capil) relatif terhadap NEXTCLOUD_SOURCE_HOME (bila di-set),
  sehingga prefix panjang tidak perlu diulang di tiap entri. Lihat resolve_source().
"""

import hashlib
import io
import json
import logging
import os
import posixpath
import re
import sys
import time
import uuid
import xml.etree.ElementTree as ET
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from urllib.parse import unquote, urlparse

import psycopg2
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from psycopg2 import sql
from psycopg2.extras import Json, execute_values

load_dotenv()

# ── Logging ────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    stream=sys.stdout,
)
log = logging.getLogger(__name__)

# ── Config ─────────────────────────────────────────────────────────────────

NEXTCLOUD_URL      = os.environ["NEXTCLOUD_URL"]           # https://cloud.example.com
NEXTCLOUD_USER     = os.environ["NEXTCLOUD_USER"]
NEXTCLOUD_PASSWORD = os.environ["NEXTCLOUD_PASSWORD"]
SCHEDULE_MINUTES   = int(os.getenv("SCHEDULE_INTERVAL_MINUTES", "60"))

# Retry saat Nextcloud membalas 423 Locked / 5xx sesaat ketika file diunduh.
WEBDAV_MAX_RETRIES           = int(os.getenv("WEBDAV_MAX_RETRIES", "5"))
WEBDAV_RETRY_BACKOFF_SECONDS = float(os.getenv("WEBDAV_RETRY_BACKOFF_SECONDS", "3"))
# (connect, read) dalam detik. Tanpa timeout, koneksi Nextcloud yang putus di
# tengah jalan membuat run menggantung selamanya — lock single-flight ikut
# tertahan, jadi jadwal berikutnya dan /load ditolak sampai container di-restart.
WEBDAV_TIMEOUT = (float(os.getenv("WEBDAV_CONNECT_TIMEOUT_SECONDS", "15")),
                  float(os.getenv("WEBDAV_READ_TIMEOUT_SECONDS", "120")))

# Warehouse tujuan. Defaultnya cocok dengan compose stack ini.
DB = {
    "host":     os.getenv("LOAD_DB_HOST", "postgres-db"),
    "port":     int(os.getenv("LOAD_DB_PORT", "5432")),
    "dbname":   os.getenv("LOAD_DB_NAME", "capil_db"),
    "user":     os.getenv("LOAD_DB_USER", "admin"),
    "password": os.getenv("LOAD_DB_PASSWORD", "password123"),
}
DB_SCHEMA = os.getenv("LOAD_DB_SCHEMA", "public")

# Zona untuk batas bulan pada aturan "lewati kalau tidak berubah" (lihat
# _muat_dataset). Sama dengan zona Laporan Rekap Bulanan supaya "bulan ini"
# berarti hal yang sama di dua tempat.
TIMEZONE = os.getenv("LOAD_TIMEZONE") or os.getenv("REPORT_TIMEZONE") or "Asia/Jakarta"

# Penanda baris yang ditulis loader ini di _airbyte_meta — membedakannya dari
# tarikan Airbyte lama, dan dipakai untuk mendeteksi Airbyte yang masih aktif.
PENANDA_LOADER = "excel-loader"

# Dataset yang dimuat container ini. Dua service, satu image: excel-loader
# (finance, perintah /load) dan capil-loader (capil, perintah /sync) — masing-
# masing punya server kontrol, lock, dan log sendiri.
LOADER_DATASET = os.getenv("LOADER_DATASET", "finance").strip().lower() or "finance"
if LOADER_DATASET not in ("finance", "capil"):
    raise ValueError(f"LOADER_DATASET={LOADER_DATASET!r} tidak dikenal (finance | capil)")
CAPIL = LOADER_DATASET == "capil"

# Capil: satu file per kantor di folder kantornya, berdampingan dengan file lain
# yang bukan urusan loader — karena itu hanya file bernama ini yang dibaca.
CAPIL_FILE  = os.getenv("CAPIL_FILE", "data.xlsx").strip() or "data.xlsx"
# Kosong = sheet pertama workbook. Diisi = sheet bernama ini (tak peka huruf besar).
CAPIL_SHEET = os.getenv("CAPIL_SHEET", "").strip()
# Tabel yang bentuknya ditiru untuk kantor capil baru; kosong = tabel capil
# pertama yang ada (raw_a1, raw_a2, …).
CAPIL_TABEL_ACUAN = os.getenv("CAPIL_TABEL_ACUAN", "").strip()

if os.getenv("NEXTCLOUD_DEST_PATH", "").strip():
    # Dulu wajib (folder hasil split). Tidak dipakai lagi; bukan error supaya
    # .env lama tetap jalan, tapi disebut supaya tidak dikira masih berpengaruh.
    print("[excel-loader] NEXTCLOUD_DEST_PATH diabaikan — tidak ada lagi file split "
          "yang di-upload.", flush=True)

# Folder induk untuk semua path sumber. Kalau di-set, entri di
# NEXTCLOUD_SOURCE_PATHS cukup ditulis relatif (mis. "A1/Finance").
SOURCE_HOME = unquote(os.getenv("NEXTCLOUD_SOURCE_HOME", "").strip()).strip("/")


def resolve_source(path: str) -> str:
    """
    Ubah satu entri path sumber menjadi path absolut di Nextcloud.

    - SOURCE_HOME kosong  -> path dipakai apa adanya (perilaku sebelum var ini ada).
    - path sudah di bawah SOURCE_HOME -> dibiarkan. Join-nya idempoten, jadi .env
      lama yang menulis prefix lengkap di tiap entri tetap jalan tanpa diedit.
    - selain itu -> digabung di bawah SOURCE_HOME.

    Hasil yang keluar dari SOURCE_HOME (mis. lewat '..') ditolak: path bisa datang
    dari argumen perintah /load di Telegram, dan itu tidak boleh bisa menunjuk ke
    folder sembarang di akun Nextcloud.
    """
    bersih = posixpath.normpath(unquote(path.strip()).replace("\\", "/")).strip("/")
    if bersih in ("", "."):
        raise ValueError("path sumber kosong")

    # normpath tidak membuang '..' yang berada di depan; tolak eksplisit. Cek ini
    # berlaku juga saat SOURCE_HOME kosong, di mana tidak ada batas lain.
    if bersih == ".." or bersih.startswith("../") or "/../" in bersih:
        raise ValueError(f"path {path!r} mengandung '..'")

    if not SOURCE_HOME:
        return f"/{bersih}"

    if bersih == SOURCE_HOME or bersih.startswith(f"{SOURCE_HOME}/"):
        hasil = bersih
    else:
        hasil = posixpath.normpath(f"{SOURCE_HOME}/{bersih}").strip("/")

    if hasil != SOURCE_HOME and not hasil.startswith(f"{SOURCE_HOME}/"):
        raise ValueError(
            f"path {path!r} keluar dari NEXTCLOUD_SOURCE_HOME ({SOURCE_HOME!r})"
        )
    return f"/{hasil}"


def parse_sources(raw: str) -> list[str]:
    """Pecah string koma/newline menjadi daftar path absolut."""
    return [resolve_source(p) for p in re.split(r"[,\n]", raw) if p.strip()]


# Dukung kedua nama variabel (NEXTCLOUD_SOURCE_PATHS dan NEXTCLOUD_SOURCE_PATH).
# Pisahkan dengan koma atau newline; unquote() handle path URL-encoded (%20, dll.).
_raw_sources = (
    os.environ.get("NEXTCLOUD_SOURCE_PATHS")
    or os.environ.get("NEXTCLOUD_SOURCE_PATH")
    or ""
)
if CAPIL:
    # Folder kantor. Kosong = induk dari setiap folder finance (A1/Finance -> A1),
    # karena data.xlsx tinggal satu tingkat di atas folder Finance.
    _raw_sources = os.environ.get("CAPIL_SOURCE_PATHS", "").strip() or ",".join(
        re.sub(r"(?i)/+finance/*$", "", p.strip()) for p in re.split(r"[,\n]", _raw_sources)
        if p.strip()
    )
    if not _raw_sources:
        raise ValueError("Set CAPIL_SOURCE_PATHS (atau NEXTCLOUD_SOURCE_PATHS) di .env")
elif not _raw_sources:
    raise ValueError("Set NEXTCLOUD_SOURCE_PATHS (atau NEXTCLOUD_SOURCE_PATH) di .env")

print(f"[excel-loader] RAW SOURCE: {repr(_raw_sources)}", flush=True)
if SOURCE_HOME:
    print(f"[excel-loader] SOURCE HOME: /{SOURCE_HOME}", flush=True)
else:
    # Tanpa SOURCE_HOME, argumen /load dari Telegram bisa menunjuk folder mana pun
    # di akun Nextcloud (selain '..', yang tetap ditolak).
    print(
        "[excel-loader] PERINGATAN: NEXTCLOUD_SOURCE_HOME kosong — "
        "argumen path pada POST /run tidak dibatasi ke satu folder induk.",
        flush=True,
    )

SOURCE_PATHS: list[str] = parse_sources(_raw_sources)

print(f"[excel-loader] PARSED {len(SOURCE_PATHS)} path(s):", flush=True)
for i, p in enumerate(SOURCE_PATHS, 1):
    print(f"[excel-loader]   [{i}] {repr(p)}", flush=True)

# Hanya ambil origin (scheme + host) dari NEXTCLOUD_URL,
# menghindari dobel path jika user menyertakan /remote.php/... di URL.
_parsed = urlparse(NEXTCLOUD_URL)
_origin = f"{_parsed.scheme}://{_parsed.netloc}"
WEBDAV_BASE = f"{_origin}/remote.php/dav/files/{NEXTCLOUD_USER}"

SESSION = requests.Session()
SESSION.auth = (NEXTCLOUD_USER, NEXTCLOUD_PASSWORD)
SESSION.headers.update({"Content-Type": "application/xml; charset=utf-8"})

# ── Kontrak dataset ─────────────────────────────────────────────────────────
# Satu entri per sheet yang dimuat. Nama kolom = header baris 1 yang sudah
# dinormalisasi (_nama_kolom: "TUNAI IFQ" -> "TUNAI_IFQ"), sama dengan nama yang
# dulu dibuat Airbyte. Tipe mengikuti tabel raw yang ada di produksi.
#
# Kolom wajib hilang = kantor itu gagal dimuat (bukan dimuat dengan NULL):
# kolom yang hilang diam-diam akan muncul di laporan sebagai angka nol.

JENIS_FINANCE = ["FI", "ZF", "AQQ", "FDY", "IFQ", "LQT", "SDQ", "SNK", "TDY", "ZKT"]

KONTRAK_FINANCE = {
    "REKAP": {
        "tabel": "raw_finance_rekap_{kantor}",
        "teks":  ["JENIS"],
        "angka": ["TOTAL_100_PERSEN", "DIKELOLA_KANWIL", "DISETOR", "PEMBULATAN_SETOR"],
    },
    "RINCIAN": {
        "tabel": "raw_finance_rincian_{kantor}",
        "teks":  ["INSTANSI"],
        "angka": [f"{p}_{j}" for p in ("TUNAI", "NOMINAL") for j in JENIS_FINANCE],
    },
}

# Capil: ±50 kolom teks yang namanya buatan Airbyte (huruf besar-kecil campur,
# mis. "Alamat_Negara", "No_" dari judul "No."). Daftarnya TIDAK ditulis di sini:
# _kontrak_capil() membacanya dari tabel raw di database, persis dengan urutan
# kolomnya. Urutan itu penting — stg_*_capil memakai dbt_utils.star() dan
# mart_capil menggabungkan semua kantor dengan UNION ALL berdasarkan posisi.
KONTRAK_CAPIL = {"CAPIL": {"tabel": "raw_{kantor}", "teks": None, "angka": []}}

DATASETS = KONTRAK_CAPIL if CAPIL else KONTRAK_FINANCE
POLA_TABEL_CAPIL = re.compile(r"^raw_[a-z][0-9]+$")

# kantor_id harus aman jadi bagian nama tabel. Pola dbt (hist_pulls) lebih
# sempit — '[a-z][0-9]+' — kantor di luar pola itu tetap dimuat tapi diberi
# peringatan karena tidak akan terbaca model riwayat.
POLA_KANTOR     = re.compile(r"^[a-z0-9_]+$")
POLA_KANTOR_DBT = re.compile(r"^[a-z][0-9]+$")

# Sheet dianggap selesai setelah sekian baris kosong total berturut-turut.
BATAS_BARIS_KOSONG = 1000

# ── WebDAV helpers ─────────────────────────────────────────────────────────

def _url(remote_path: str) -> str:
    """
    Bangun URL WebDAV lengkap.
    Toleran terhadap path yang diawali /{username}/ (Nextcloud kadang
    menampilkan path dalam format tersebut di UI) — prefix username
    dibuang karena WEBDAV_BASE sudah menyertakannya.
    """
    normalized = remote_path.lstrip("/")
    user_prefix = f"{NEXTCLOUD_USER}/"
    if normalized.startswith(user_prefix):
        normalized = normalized[len(user_prefix):]
    return f"{WEBDAV_BASE}/{normalized}"


def list_xlsx(folder: str) -> list[str]:
    """Kembalikan daftar nama file .xlsx di folder WebDAV (tidak rekursif)."""
    resp = SESSION.request(
        "PROPFIND",
        _url(folder),
        headers={"Depth": "1"},
        timeout=WEBDAV_TIMEOUT,
        data=b"""<?xml version="1.0" encoding="utf-8"?>
<d:propfind xmlns:d="DAV:">
  <d:prop><d:displayname/><d:resourcetype/></d:prop>
</d:propfind>""",
    )
    resp.raise_for_status()

    ns = {"d": "DAV:"}
    tree = ET.fromstring(resp.content)

    # Gunakan _url() agar normalisasi username-prefix konsisten
    folder_href = urlparse(_url(folder)).path.rstrip("/")

    results = []
    for node in tree.findall(".//d:response", ns):
        href = (node.findtext("d:href", namespaces=ns) or "").rstrip("/")
        # href sudah URL-encoded (mis. %20); decode agar cocok dengan nama asli
        name = unquote(href.split("/")[-1])
        # Lewati folder itu sendiri
        if href.endswith(folder_href):
            continue
        # Lewati direktori (resourcetype berisi <d:collection/>) — hanya file yang diproses
        if node.find(".//d:resourcetype/d:collection", ns) is not None:
            continue
        # HANYA file .xlsx asli — abaikan format lain (.xls, .csv, .pdf, dst.)
        if not name.lower().endswith(".xlsx"):
            log.info("    ⏭ Lewati (bukan .xlsx): %s", name)
            continue
        # Abaikan file lock/temp Office (mis. ~$laporan.xlsx) yang bukan workbook valid
        if name.startswith("~$") or name.startswith("."):
            log.info("    ⏭ Lewati (file temp/lock): %s", name)
            continue
        results.append(name)
    return results


def download(folder: str, filename: str) -> bytes:
    """
    Unduh satu file. Diulang saat 423 Locked atau 5xx dengan backoff linear —
    keduanya biasanya sesaat (OnlyOffice sedang menyimpan, Nextcloud sibuk).
    """
    url = _url(str(PurePosixPath(folder) / filename))
    resp = None
    for attempt in range(1, WEBDAV_MAX_RETRIES + 1):
        try:
            resp = SESSION.get(url, timeout=WEBDAV_TIMEOUT)
            if resp.status_code != 423 and resp.status_code < 500:
                break
            sebab = str(resp.status_code)
        except (requests.ConnectionError, requests.Timeout) as exc:
            # Putus/timeout juga biasanya sesaat; percobaan terakhir dilempar ke atas.
            if attempt == WEBDAV_MAX_RETRIES:
                raise
            sebab = type(exc).__name__
        if attempt < WEBDAV_MAX_RETRIES:
            wait = WEBDAV_RETRY_BACKOFF_SECONDS * attempt
            log.warning("    ⏳ %s saat unduh, percobaan %d/%d — tunggu %.0fs: %s",
                        sebab, attempt, WEBDAV_MAX_RETRIES, wait, filename)
            time.sleep(wait)
    resp.raise_for_status()
    return resp.content


# ── Baca & validasi sheet ───────────────────────────────────────────────────

def _nama_kolom(header) -> str:
    """'TOTAL 100 PERSEN' -> 'TOTAL_100_PERSEN' (sama dengan normalisasi Airbyte)."""
    return re.sub(r"\W+", "_", str(header).strip()).strip("_").upper()


def _lipat(nama) -> str:
    """
    Kunci pencocokan judul kolom dengan nama kolom kontrak, tak peka huruf besar
    dan tanda baca: 'No.' / 'No_' -> 'no', 'TUNAI IFQ' / 'TUNAI_IFQ' -> 'tunai_ifq'.
    Yang ditulis ke database tetap nama kanonik dari kontrak.
    """
    # Garis bawah ikut dilipat: Airbyte menamai 'Aset Hak Milik Tanah (m2)' sebagai
    # 'Aset_Hak_Milik_Tanah__m2_' (dobel + ekor '_'), dan itu harus tetap cocok.
    return re.sub(r"[\W_]+", "_", str(nama).strip()).strip("_").lower()


def _ke_angka(nilai):
    """Nilai sel -> Decimal, None untuk sel kosong. ValueError kalau bukan angka."""
    if nilai is None:
        return None
    if isinstance(nilai, bool) or isinstance(nilai, (datetime, date)):
        raise ValueError(f"bukan angka: {nilai!r}")
    if isinstance(nilai, (int, Decimal)):
        return Decimal(nilai)
    if isinstance(nilai, float):
        # str() = representasi terpendek, sama dengan yang dulu lewat JSON Airbyte.
        return Decimal(str(nilai))
    teks = str(nilai).strip()
    if not teks:
        return None
    try:
        return Decimal(teks)
    except InvalidOperation:
        raise ValueError(f"bukan angka: {teks!r}") from None


def _ke_teks(nilai):
    if nilai is None:
        return None
    if isinstance(nilai, float) and nilai.is_integer():
        nilai = int(nilai)  # 123.0 dari Excel -> '123', bukan '123.0'
    teks = str(nilai)
    return teks if teks.strip() else None


# ── Masalah data ────────────────────────────────────────────────────────────
# Setiap temuan dicatat terstruktur (bukan kalimat jadi) supaya relay bisa
# menyusun SATU rekap berbahasa sederhana untuk semua kantor — lihat
# scripts/rekap_masalah.py, yang juga memegang kalimat untuk tiap jenis.
#
# galat=True  -> kantor itu tidak dimuat (tarikan sebelumnya tetap dipakai)
# galat=False -> catatan saja, data tetap dimuat
# sistem=True -> bukan kesalahan pengisi file (Nextcloud/database bermasalah)

BATAS_MASALAH_PER_KANTOR = 200  # /status ikut membawa daftar ini tiap poll relay


def _masalah(jenis: str, galat: bool = True, **rinci) -> dict:
    return {"jenis": jenis, "galat": galat, **{k: v for k, v in rinci.items() if v is not None}}


def _nilai_pendek(nilai) -> str:
    teks = str(nilai)
    return teks if len(teks) <= 40 else teks[:37] + "..."


def baca_sheet(ws_nilai, ws_rumus, kontrak: dict) -> tuple[list[dict], list[dict]]:
    """
    Ubah satu sheet menjadi daftar baris sesuai kontrak.

    ws_nilai dibuka dengan data_only=True (hasil hitung tersimpan), ws_rumus
    tanpa — dipakai hanya untuk mengenali sel formula yang tidak punya hasil
    tersimpan, yang di ws_nilai tampak kosong.

    Kembalikan (baris, masalah). Pemeriksaan jalan terus sampai akhir sheet
    supaya SEMUA masalah terlihat sekaligus — pengisi file cukup memperbaiki
    sekali, bukan satu per satu tiap kali loader jalan. Satu masalah galat saja
    sudah membatalkan kantor itu, karena dimuat sebagian lebih buruk dari tidak
    dimuat. Airbyte dulu memuat formula tanpa hasil sebagai kosong tanpa suara,
    dan tarikan bertotal kosong itu menjadi "tarikan terakhir" — laporan jadi nol.
    """
    masalah: list[dict] = []
    it_nilai = ws_nilai.iter_rows(values_only=True)
    it_rumus = ws_rumus.iter_rows(values_only=True)

    header = next(it_nilai, None)
    next(it_rumus, None)
    if not header or all(h is None or not str(h).strip() for h in header):
        return [], [_masalah("sheet_kosong")]

    wajib = kontrak["teks"] + kontrak["angka"]
    kanonik = {_lipat(k): k for k in wajib}

    posisi: dict[str, int] = {}
    for i, h in enumerate(header):
        if h is None or not str(h).strip():
            continue
        # Judul di luar kontrak: capil ditulis seperti di file (namanya campur
        # huruf besar-kecil), finance dinormalisasi seperti sebelumnya.
        nama = kanonik.get(_lipat(h)) or (str(h).strip() if CAPIL else _nama_kolom(h))
        if nama in posisi:
            masalah.append(_masalah("kolom_ganda", kolom=nama,
                                    sel=f"{get_column_letter(i + 1)}1"))
            continue
        posisi[nama] = i

    for k in wajib:
        if k not in posisi:
            masalah.append(_masalah("kolom_hilang", kolom=k))
    tambahan = [k for k in posisi if k not in wajib]
    if tambahan:
        masalah.append(_masalah("kolom_tambahan", galat=False, kolom=", ".join(tambahan)))
    ada = [k for k in wajib if k in posisi]  # kolom yang hilang tetap dilaporkan di atas

    baris: list[dict] = []
    rumus_kosong: dict[str, list[str]] = {}
    kosong_beruntun = 0
    for nomor, (nilai, rumus) in enumerate(zip(it_nilai, it_rumus), start=2):
        # Template finance memformat seluruh kolom sampai baris 1.048.576, dan
        # mode read_only ikut menelusuri semuanya (~4 detik per sheet). Deretan
        # panjang baris kosong total = akhir data.
        if all(v is None for v in nilai) and all(v is None for v in rumus):
            kosong_beruntun += 1
            if kosong_beruntun >= BATAS_BARIS_KOSONG:
                break
            continue
        kosong_beruntun = 0

        def sel(kolom):
            i = posisi[kolom]
            return nilai[i] if i < len(nilai) else None

        mentah = {k: sel(k) for k in ada}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in mentah.values()):
            continue  # baris kosong (ekor sheet, baris pemisah)

        rekaman = {k: None for k in wajib}
        for k in kontrak["teks"]:
            if k in mentah:
                rekaman[k] = _ke_teks(mentah[k])
        for k in kontrak["angka"]:
            if k not in mentah:
                continue
            alamat = f"{get_column_letter(posisi[k] + 1)}{nomor}"
            try:
                rekaman[k] = _ke_angka(mentah[k])
            except ValueError:
                masalah.append(_masalah("bukan_angka", sel=alamat, kolom=k,
                                        nilai=_nilai_pendek(mentah[k])))
        for k in ada:
            i = posisi[k]
            f = rumus[i] if i < len(rumus) else None
            if mentah[k] is None and isinstance(f, str) and f.startswith("="):
                alamat = f"{get_column_letter(i + 1)}{nomor}"
                if CAPIL:
                    # Capil: catatan saja, dikumpulkan per kolom. Kolom formulanya
                    # (Usia_*, Status_Pemberdayaan, ...) turunan dan tidak dipakai
                    # dbt/laporan; Airbyte dulu juga memuatnya kosong.
                    rumus_kosong.setdefault(k, []).append(alamat)
                else:
                    # Finance: galat — total kosong jadi nol di laporan.
                    masalah.append(_masalah("formula_tanpa_hasil", sel=alamat, kolom=k))
        baris.append(rekaman)

    for k, sel_sel in rumus_kosong.items():
        masalah.append(_masalah("formula_tanpa_hasil", galat=False, kolom=k,
                                sel=sel_sel[0], jumlah=len(sel_sel)))

    return baris, masalah


def _sidik(baris: list[dict]) -> str:
    """Sidik isi data (bukan byte file): simpan ulang tanpa perubahan = sidik sama."""
    kanonik = json.dumps(baris, sort_keys=True, default=str, ensure_ascii=False)
    return hashlib.sha256(kanonik.encode("utf-8")).hexdigest()


def _kunci_sheet(nama: str, urutan: int):
    """Sheet ini dimuat sebagai dataset apa? None = bukan urusan kontrak."""
    if not CAPIL:
        kunci = nama.strip().upper()
        return kunci if kunci in DATASETS else None
    if CAPIL_SHEET:
        return "CAPIL" if nama.strip().upper() == CAPIL_SHEET.upper() else None
    return "CAPIL" if urutan == 0 else None


def nama_sheet(kunci: str) -> str:
    """Nama sheet yang diharapkan untuk satu kunci dataset (untuk pesan sheet_hilang)."""
    return CAPIL_SHEET if CAPIL and CAPIL_SHEET else kunci


def baca_workbook(isi: bytes, kontrak: dict) -> tuple[dict, list[dict]]:
    """
    Baca semua sheet yang dikenal kontrak ({kunci: kontrak sheet}).
    Kembalikan ({KUNCI: baris}, masalah) — setiap masalah sudah membawa
    nama sheet-nya. Sheet lain diabaikan dengan catatan. File yang tidak bisa
    dibuka sama sekali melempar exception.
    """
    wb_nilai = load_workbook(io.BytesIO(isi), data_only=True, read_only=True)
    wb_rumus = load_workbook(io.BytesIO(isi), data_only=False, read_only=True)
    hasil: dict[str, list[dict]] = {}
    masalah: list[dict] = []
    try:
        for urutan, nama in enumerate(wb_nilai.sheetnames):
            kunci = _kunci_sheet(nama, urutan)
            if kunci is None:
                masalah.append(_masalah("sheet_tak_dikenal", galat=False, sheet=nama))
                continue
            if kunci in hasil:
                masalah.append(_masalah("sheet_ganda", sheet=nama))
                continue
            baris, temuan = baca_sheet(wb_nilai[nama], wb_rumus[nama], kontrak[kunci])
            hasil[kunci] = baris
            masalah += [{**m, "sheet": nama} for m in temuan]
    finally:
        wb_nilai.close()
        wb_rumus.close()
    return hasil, masalah


# ── Tulis ke Postgres ───────────────────────────────────────────────────────

def _pastikan_tabel(cur, tabel: str, kontrak: dict) -> None:
    """
    Buat tabel raw kalau belum ada, dengan bentuk yang sama dengan buatan Airbyte.
    Kantor baru jadi tidak butuh koneksi Airbyte — cukup foldernya di
    NEXTCLOUD_SOURCE_PATHS / CAPIL_SOURCE_PATHS dan tabelnya di dbt sources.yml.

    Kontrak capil membawa `acuan`: tabel kantor lain yang ditiru utuh (LIKE),
    supaya urutan kolomnya sama dan UNION ALL di mart_capil tetap cocok.
    """
    ident = sql.Identifier(DB_SCHEMA, tabel)
    if kontrak.get("acuan"):
        cur.execute(sql.SQL("CREATE TABLE IF NOT EXISTS {} (LIKE {})").format(
            ident, sql.Identifier(DB_SCHEMA, kontrak["acuan"])))
        return
    kolom = [sql.SQL("{} character varying").format(sql.Identifier(k)) for k in kontrak["teks"]]
    kolom += [sql.SQL("{} numeric").format(sql.Identifier(k)) for k in kontrak["angka"]]
    cur.execute(sql.SQL("""
        CREATE TABLE IF NOT EXISTS {} (
            _airbyte_raw_id character varying NOT NULL,
            _airbyte_extracted_at timestamp with time zone NOT NULL,
            _airbyte_meta jsonb NOT NULL,
            _airbyte_generation_id bigint NOT NULL,
            {}
        )""").format(ident, sql.SQL(", ").join(kolom)))
    # Tabel lama buatan Airbyte bisa kekurangan kolom yang nilainya selalu kosong
    # saat sync pertamanya; tambahkan supaya INSERT tidak gagal.
    for k in kontrak["teks"]:
        cur.execute(sql.SQL("ALTER TABLE {} ADD COLUMN IF NOT EXISTS {} character varying")
                    .format(ident, sql.Identifier(k)))
    for k in kontrak["angka"]:
        cur.execute(sql.SQL("ALTER TABLE {} ADD COLUMN IF NOT EXISTS {} numeric")
                    .format(ident, sql.Identifier(k)))


def _muat_dataset(cur, tabel: str, kontrak: dict, baris: list[dict], meta: dict) -> dict:
    """
    Tambahkan satu tarikan ke tabel raw. Kembalikan ringkasan untuk laporan.

    Dilewati bila isi datanya sama dengan tarikan terakhir DAN tarikan itu sudah
    terjadi di bulan berjalan: run per jam tidak lagi menggandakan data yang
    sama (masalah mode Append Airbyte), tapi setiap bulan tetap punya minimal
    satu tarikan sendiri, sehingga potret Rekap Bulanan tidak perlu "meminjam"
    tarikan bulan sebelumnya.
    """
    ident = sql.Identifier(DB_SCHEMA, tabel)
    _pastikan_tabel(cur, tabel, kontrak)
    # Satu penulis per tabel sampai commit — generation id dihitung dari MAX.
    cur.execute(sql.SQL("LOCK TABLE {} IN SHARE ROW EXCLUSIVE MODE").format(ident))

    cur.execute(sql.SQL("""
        WITH akhir AS (SELECT MAX(_airbyte_generation_id) AS gen FROM {t})
        SELECT a.gen,
               (SELECT r._airbyte_meta->>'sha256' FROM {t} r
                 WHERE r._airbyte_generation_id = a.gen LIMIT 1),
               (SELECT date_trunc('month', MIN(r._airbyte_extracted_at) AT TIME ZONE %(tz)s)
                     = date_trunc('month', now() AT TIME ZONE %(tz)s)
                  FROM {t} r WHERE r._airbyte_generation_id = a.gen),
               (SELECT COUNT(*) FROM {t} r
                 WHERE r._airbyte_meta->>'loader' IS NULL
                   AND r._airbyte_extracted_at > (
                       SELECT MAX(x._airbyte_extracted_at) FROM {t} x
                        WHERE x._airbyte_meta->>'loader' = %(penanda)s))
        FROM akhir a""").format(t=ident), {"tz": TIMEZONE, "penanda": PENANDA_LOADER})
    gen_akhir, sidik_akhir, bulan_sama, baris_airbyte = cur.fetchone()

    ringkasan = {"tabel": tabel, "baris": len(baris)}
    if baris_airbyte:
        # Airbyte masih menulis setelah loader mengambil alih: generation id-nya
        # bisa lebih kecil dari milik loader dan tarikannya jadi tak terlihat.
        ringkasan["airbyte_aktif"] = True
        log.warning("    ⚠ %s: %d baris baru dari Airbyte setelah muatan loader — "
                    "matikan koneksi Airbyte %s kantor ini.", tabel, baris_airbyte, LOADER_DATASET)

    if sidik_akhir == meta["sha256"] and bulan_sama:
        ringkasan["status"] = "tidak berubah"
        ringkasan["tarikan"] = gen_akhir
        return ringkasan

    gen = (gen_akhir or 0) + 1
    kolom = kontrak["teks"] + kontrak["angka"]
    nilai = [
        (str(uuid.uuid4()), Json(meta), gen, *[r[k] for k in kolom])
        for r in baris
    ]
    if nilai:
        execute_values(
            cur,
            sql.SQL("INSERT INTO {} (_airbyte_raw_id, _airbyte_extracted_at, _airbyte_meta, "
                    "_airbyte_generation_id, {}) VALUES %s").format(
                ident, sql.SQL(", ").join(map(sql.Identifier, kolom))
            ).as_string(cur),
            nilai,
            # now() = waktu mulai transaksi: sama untuk semua baris satu tarikan,
            # jadi ditarik_pada di hist_pulls tidak terpecah.
            template="(%s, now(), %s, %s" + ", %s" * len(kolom) + ")",
        )
    ringkasan["status"] = "dimuat"
    ringkasan["tarikan"] = gen
    return ringkasan


# Tipe kolom raw yang diperlakukan sebagai angka; selain ini dibaca sebagai teks.
TIPE_ANGKA = {"numeric", "bigint", "integer", "smallint", "double precision", "real"}


def _kolom_tabel(cur, tabel: str) -> list[tuple[str, str]]:
    """Kolom data (tanpa _airbyte_*) suatu tabel, urut posisi. Kosong = tabel tidak ada."""
    cur.execute("""
        SELECT column_name, data_type FROM information_schema.columns
         WHERE table_schema = %s AND table_name = %s
           AND column_name NOT LIKE '\\_airbyte\\_%%'
         ORDER BY ordinal_position""", (DB_SCHEMA, tabel))
    return cur.fetchall()


def _kontrak_capil(conn, tabel: str) -> dict:
    """
    Kontrak sheet capil untuk satu kantor, dibaca dari database.

    Tabel kantor itu sudah ada (dibuat Airbyte dulu) -> kolomnya itulah kontraknya,
    persis nama dan urutannya. Belum ada (kantor baru) -> tiru tabel capil kantor
    lain (CAPIL_TABEL_ACUAN, atau yang pertama ditemukan). Melempar LookupError
    kalau tidak ada satu pun tabel capil untuk ditiru.
    """
    with conn, conn.cursor() as cur:
        kolom, acuan = _kolom_tabel(cur, tabel), None
        if not kolom:
            if CAPIL_TABEL_ACUAN:
                calon = [CAPIL_TABEL_ACUAN]
            else:
                cur.execute("SELECT table_name FROM information_schema.tables "
                            "WHERE table_schema = %s ORDER BY table_name", (DB_SCHEMA,))
                calon = [t for (t,) in cur.fetchall() if POLA_TABEL_CAPIL.match(t)]
            for t in calon:
                kolom = _kolom_tabel(cur, t)
                if kolom:
                    acuan = t
                    break
    # Kolom 'null' = buatan Airbyte untuk judul kolom yang kosong; tidak bisa
    # dicocokkan dengan judul apa pun, jadi bukan kontrak — dibiarkan NULL.
    kolom = [(k, t) for k, t in kolom if k.lower() != "null"]
    if not kolom:
        raise LookupError(f"tabel {tabel} belum ada dan tidak ada tabel capil lain "
                          f"di schema {DB_SCHEMA} untuk ditiru bentuknya")
    return {"CAPIL": {
        "tabel": KONTRAK_CAPIL["CAPIL"]["tabel"],
        "teks":  [k for k, tipe in kolom if tipe not in TIPE_ANGKA],
        "angka": [k for k, tipe in kolom if tipe in TIPE_ANGKA],
        "acuan": acuan,
    }}


# ── Proses utama ────────────────────────────────────────────────────────────

def _kantor_name(source_path: str) -> str:
    """
    Ambil nama kantor dari path folder sumber.

    Struktur folder: <nama kantor>/Finance/finance.xlsx — folder yang di-scan
    adalah <nama kantor>/Finance, sehingga nama kantor = komponen tepat DI ATAS
    subfolder 'Finance'. Jika komponen terakhir bukan 'Finance', komponen
    terakhir dipakai apa adanya — itulah kasus capil (<nama kantor>/data.xlsx).
    """
    parts = [p for p in source_path.strip("/").split("/") if p]
    if not parts:
        return "unknown"
    if len(parts) >= 2 and parts[-1].lower() == "finance":
        return parts[-2]
    return parts[-1]


def _log_masalah(m: dict) -> None:
    lokasi = " ".join(f"{k}={m[k]}" for k in ("file", "sheet", "sel", "kolom", "nilai") if k in m)
    (log.error if m["galat"] else log.warning)(
        "    %s %s %s%s", "✗" if m["galat"] else "⚠", m["jenis"], lokasi,
        f" — {m['detail']}" if m.get("detail") else "")


def _tolak(hasil: dict) -> dict:
    """Tandai kantor ditolak; `error` = ringkasan pendek untuk log & Mini App."""
    galat = [m for m in hasil["masalah"] if m["galat"]]
    if len(hasil["masalah"]) > BATAS_MASALAH_PER_KANTOR:
        sisa = len(hasil["masalah"]) - BATAS_MASALAH_PER_KANTOR
        hasil["masalah"] = hasil["masalah"][:BATAS_MASALAH_PER_KANTOR] + [
            _masalah("terpotong", jumlah=sisa)]
    hasil["error"] = f"{len(galat)} masalah — data kantor ini tidak dimuat"
    log.error("  ✗ %s: %s", hasil["kantor"], hasil["error"])
    return hasil


def process_source(conn, source_path: str) -> dict:
    """
    Muat semua sheet dataset ini dari satu folder sumber dalam SATU transaksi.

    Semua-atau-tidak-sama-sekali per kantor: kalau satu sheet bermasalah atau
    gagal ditulis, tidak ada tabel kantor itu yang berubah — sama dengan aturan
    split-excel dulu ("satu sheet gagal, seluruh kantor dilewati"). Semua file
    dan sheet tetap diperiksa sampai habis supaya rekapnya lengkap.

    Kembalikan ringkasan: {kantor, ok, tabel: [...], masalah: [...], error?}.
    """
    nama_kantor = _kantor_name(source_path)
    kantor_id = nama_kantor.lower()
    hasil = {"kantor": nama_kantor.upper(), "sumber": source_path, "ok": False,
             "tabel": [], "masalah": []}
    catat = hasil["masalah"].append
    log.info("── Sumber: %s  (kantor: %s)", source_path, kantor_id)

    if not POLA_KANTOR.match(kantor_id):
        catat(_masalah("nama_kantor_tidak_valid", sistem=True, detail=nama_kantor))
        return _tolak(hasil)
    if not POLA_KANTOR_DBT.match(kantor_id):
        catat(_masalah("kantor_di_luar_pola", galat=False, sistem=True, detail=kantor_id))

    try:
        files = list_xlsx(source_path)
    except Exception as exc:
        catat(_masalah("folder_gagal", sistem=True, detail=str(exc)[:300]))
        return _tolak(hasil)

    if CAPIL:
        # Folder kantor berisi file lain juga; yang dibaca hanya CAPIL_FILE.
        cocok = [f for f in files if f.lower() == CAPIL_FILE.lower()]
        if not cocok:
            catat(_masalah("file_hilang", file=CAPIL_FILE,
                           detail=", ".join(files[:10]) or None))
            return _tolak(hasil)
        files = cocok[:1]
        tabel = DATASETS["CAPIL"]["tabel"].format(kantor=kantor_id)
        try:
            kontrak = _kontrak_capil(conn, tabel)
        except Exception as exc:
            catat(_masalah("kontrak_gagal", sistem=True, detail=str(exc)[:300]))
            return _tolak(hasil)
        if kontrak["CAPIL"]["acuan"]:
            log.info("  %s belum ada — dibuat meniru %s.", tabel, kontrak["CAPIL"]["acuan"])
    else:
        kontrak = DATASETS
        if not files:
            catat(_masalah("folder_kosong"))
            return _tolak(hasil)
    log.info("  Ditemukan %d file: %s", len(files), files)

    # Baca & periksa semuanya dulu, baru buka transaksi.
    data: dict[str, tuple[list[dict], dict]] = {}
    ada_file_rusak = False
    for filename in files:
        log.info("  Membaca: %s", filename)
        try:
            isi = download(source_path, filename)
        except Exception as exc:
            catat(_masalah("unduh_gagal", sistem=True, file=filename, detail=str(exc)[:300]))
            ada_file_rusak = True
            continue
        try:
            sheets, temuan = baca_workbook(isi, kontrak)
        except Exception as exc:
            log.error("    ✗ Gagal membuka %s: %s: %s", filename, type(exc).__name__, exc)
            catat(_masalah("file_rusak", file=filename, detail=str(exc)[:300]))
            ada_file_rusak = True
            continue
        hasil["masalah"] += [{**m, "file": filename} for m in temuan]
        for kunci, baris in sheets.items():
            if kunci in data:
                catat(_masalah("sheet_ganda_antar_file", file=filename, sheet=kunci,
                               detail=data[kunci][1]["source_file"].rsplit("/", 1)[-1]))
                continue
            data[kunci] = (baris, {
                "loader": PENANDA_LOADER,
                "source_file": f"{source_path}/{filename}",
                "sha256": _sidik(baris),
            })
            log.info("    %s: %d baris", kunci, len(baris))

    # Sheet yang hilang baru bermakna kalau semua file bisa dibuka.
    if not ada_file_rusak:
        for kunci in DATASETS:
            if kunci not in data:
                catat(_masalah("sheet_hilang", sheet=nama_sheet(kunci)))

    for m in hasil["masalah"]:
        _log_masalah(m)
    if any(m["galat"] for m in hasil["masalah"]):
        return _tolak(hasil)

    try:
        with conn:  # commit bila sukses, rollback bila ada exception
            with conn.cursor() as cur:
                for kunci, (baris, meta) in data.items():
                    tabel = kontrak[kunci]["tabel"].format(kantor=kantor_id)
                    r = _muat_dataset(cur, tabel, kontrak[kunci], baris, meta)
                    hasil["tabel"].append(r)
                    log.info("    %s %s: %s (%d baris, tarikan #%s)",
                             "✓" if r["status"] == "dimuat" else "=",
                             tabel, r["status"], r["baris"], r["tarikan"])
    except Exception as exc:
        log.error("  ✗ Gagal menulis — tidak ada tabel kantor ini yang berubah.", exc_info=True)
        hasil["tabel"] = []
        catat(_masalah("gagal_simpan", sistem=True, detail=str(exc)[:300]))
        return _tolak(hasil)

    hasil["ok"] = True
    return hasil


def run_once(sources: list[str] | None = None) -> dict:
    """
    Jalankan satu siklus. Kembalikan {"ok": bool, "kantor": [ringkasan per kantor]} —
    job_control menyimpannya sebagai last_result, relay menyusun pesan Telegram
    darinya.

    `sources` None = pakai SOURCE_PATHS dari .env (perilaku terjadwal). Daftar
    eksplisit dipakai saat run dipicu manual, mis. `/load A1/Finance` dari
    Telegram — path-nya sudah lewat resolve_source() di validate_params().
    """
    targets = sources or SOURCE_PATHS

    log.info("=== Mulai run excel-loader (%s) ===", LOADER_DATASET)
    log.info("  Tujuan  : %s@%s/%s (schema %s)", DB["user"], DB["host"], DB["dbname"], DB_SCHEMA)
    log.info("  Sumber  : %d folder%s", len(targets), " (dipilih manual)" if sources else "")
    for i, p in enumerate(targets, 1):
        log.info("    [%d] %s", i, p)

    try:
        conn = psycopg2.connect(connect_timeout=15, **DB)
    except Exception as exc:
        log.error("Gagal konek ke database: %s", exc)
        return {"ok": False, "dataset": LOADER_DATASET,
                "error": f"gagal konek ke database: {str(exc)[:300]}", "kantor": []}

    try:
        ringkasan = [process_source(conn, s) for s in targets]
    finally:
        conn.close()

    ok = all(r["ok"] for r in ringkasan)
    if ok:
        log.info("=== Run selesai (semua sukses) ===")
    else:
        log.error("=== Run selesai DENGAN KEGAGALAN — lihat error di atas ===")
    return {"ok": ok, "dataset": LOADER_DATASET,
            # Nama sheet yang dibaca, untuk kalimat rekap; [] = sheet pertama.
            "sheet_dibaca": [nama_sheet(k) for k in DATASETS] if not CAPIL or CAPIL_SHEET else [],
            "kantor": ringkasan}


def validate_params(params: dict) -> dict:
    """
    Validasi body POST /run. Melempar ValueError -> dibalas 400 oleh job_control.

    Ini satu-satunya tempat path dari luar diperiksa; bot Telegram sengaja tidak
    ikut memvalidasi supaya tidak ada dua aturan yang bisa berbeda.
    """
    raw = params.get("sources")
    if raw in (None, "", []):
        return {}
    if isinstance(raw, str):
        raw = re.split(r"[,\s]+", raw)
    if not isinstance(raw, list):
        raise ValueError("field 'sources' harus berupa list atau string")

    sources = [resolve_source(str(p)) for p in raw if str(p).strip()]
    if not sources:
        raise ValueError("field 'sources' kosong setelah dibersihkan")
    return {"sources": sources}


def main() -> None:
    flag = {a for a in sys.argv[1:] if a.startswith("-")}
    tak_dikenal = flag - {"--watch", "--serve"}
    if tak_dikenal:
        log.error("Opsi tidak dikenal: %s (yang ada: --watch, --serve)",
                  ", ".join(sorted(tak_dikenal)))
        sys.exit(2)

    # Argumen non-flag = daftar folder sumber, seperti argumen /load di Telegram.
    argumen = [a for a in sys.argv[1:] if not a.startswith("-")]
    pilihan = None
    if argumen:
        if "--watch" in flag:
            log.error("Argumen folder tidak bisa digabung dengan --watch.")
            sys.exit(2)
        try:
            pilihan = validate_params({"sources": argumen}).get("sources")
        except ValueError as exc:
            log.error("Argumen ditolak: %s", exc)
            sys.exit(2)

    runner = None
    if "--serve" in flag:
        import job_control

        runner = job_control.serve(
            # Nama job = perintah bot yang memicunya: /load (finance), /sync (capil).
            "sync" if CAPIL else "load",
            lambda params: run_once(params.get("sources")),
            validate=validate_params,
            background="--watch" in flag,
            info={
                "source_home": f"/{SOURCE_HOME}" if SOURCE_HOME else None,
                "sources": SOURCE_PATHS,
                "db": f"{DB['host']}/{DB['dbname']}.{DB_SCHEMA}",
                "interval_menit": SCHEDULE_MINUTES,
                "fitur": "load-db",  # penanda image loader (bukan split-excel lama)
                "dataset": LOADER_DATASET,
                **({"file": CAPIL_FILE, "sheet": CAPIL_SHEET or "(sheet pertama)"}
                   if CAPIL else {}),
            },
        )
        if "--watch" not in flag:
            return

    if "--watch" in flag:
        log.info("Mode watch aktif, interval: %d menit", SCHEDULE_MINUTES)
        while True:
            if runner is None:
                run_once()
            elif not runner.run_now({}, trigger="jadwal"):
                # Run manual dari Telegram sedang jalan — dua run bersamaan akan
                # berebut generation id, jadi jadwal ini dilewati saja.
                log.info("Run terjadwal dilewati: run lain masih berjalan.")
            log.info("Tunggu %d menit...", SCHEDULE_MINUTES)
            time.sleep(SCHEDULE_MINUTES * 60)
    else:
        sys.exit(0 if run_once(pilihan)["ok"] else 1)


if __name__ == "__main__":
    main()
