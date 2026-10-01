#!/usr/bin/env python3
"""
report_finance.py — Laporan Keuangan Bulanan (Rekapitulasi Simpanan Bulanan)

Mengisi template `docs/template/monthly-finance-report.xlsx`: tabel atas berisi
satu baris per kantor untuk bulan laporan, tabel bawah berisi total semua kantor
per bulan (Januari s.d. bulan laporan, tahun yang sama).

Dipakai oleh notif_relay.py: perintah /keuangan dan tombol "Laporan keuangan"
di Mini App memanggil bangun_laporan(), lalu hasilnya dikirim ke Telegram
sebagai dokumen — sama seperti Rekap Bulanan (report_summary.py), dan dengan
cakupan data yang sama: potret akhir bulan dari view `hist_*`, bukan staging.

Isi kolom (per kantor, per bulan). Kolom dicari dari judulnya di baris header
(HEADER), bukan dari hurufnya — template yang kolomnya digeser tetap terisi
benar, dan judul yang hilang/berubah menggagalkan laporan dengan pesan jelas:
  NASABAH TUNAI        SUM(rincian TUNAI_IFQ)
  NASABAH AKTIF 1/2/3  jumlah baris capil dengan K = 1/2/3
  TOTAL ANGGOTA        aktif 1 + 2 + 3
  SIMPANAN WAJIB       rekap NOMINAL IFQ · TOTAL 100%
  SIMPANAN POKOK       rekap NOMINAL ZKT · TOTAL 100%
  SIMPANAN SUKARELA    rekap NOMINAL SDQ · TOTAL 100%
  TOTAL KONTRIBUSI     wajib + pokok + sukarela
  TOTAL SETOR          SUM(rekap DISETOR) semua baris kecuali TOTAL
  CAD BULAN INI        rekap CAD · DISETOR
  AKUMULASI CAD        jumlah CAD semua bulan s.d. bulan laporan, hanya
                       bulan yang punya tarikan finance sendiri
  JUMLAH PETUGAS       capil dengan LMG bukan PRA / PJ% / KPJ% (= pengurus)

Grafik (sheet GRAFIK) dibuat oleh kode, BUKAN disimpan di template: openpyxl
membuang chart yang sudah ada saat membuka workbook, jadi chart di template
akan hilang dari setiap laporan. Semua grafik merujuk sel tabel laporan itu
sendiri, sehingga tetap benar kalau angkanya diedit tangan di Excel.

AKUMULASI CAD dihitung ulang dari warehouse setiap kali, mulai dari bulan
tarikan finance pertama. Jadi ia hanya selengkap riwayat tarikan finance
(Append) — kalau tabel raw finance pernah di-reset, bulan-bulan sebelumnya
hilang dari akumulasi.

Konfigurasi: sama dengan report_summary.py (REPORT_DB_*, REPORT_SCHEMA,
REPORT_VIEW_HIST_*, REPORT_TIMEZONE, REPORT_XLSX_PASSWORD), ditambah
  REPORT_FINANCE_TEMPLATE  (default docs/template/monthly-finance-report.xlsx)
"""

import io
import logging
import os
from datetime import date
from pathlib import Path

import report_summary as rs
from report_summary import ReportError

try:
    from openpyxl.chart import BarChart, LineChart, Reference, Series
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Font
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:  # pragma: no cover -- laporan_aktif() sudah mematikan fiturnya
    BarChart = None

log = logging.getLogger("notif_relay.report_finance")

TEMPLATE = Path(os.getenv(
    "REPORT_FINANCE_TEMPLATE",
    Path(__file__).resolve().parent.parent / "docs" / "template" / "monthly-finance-report.xlsx",
))

NAMA_BULAN = ["", "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
              "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER"]


def _pilih(tarikan: str) -> str:
    """
    Untuk tiap bulan di `bulan` dan tiap kantor: tarikan terakhir (`tarikan`
    berisi kantor_id, tarikan_id, ditarik_pada) sebelum awal bulan berikutnya,
    jam lokal — aturan potret yang sama dengan report_summary._potret().
    """
    return f"""
    SELECT b.awal_bulan, x.kantor_id, MAX(x.tarikan_id) AS tarikan_id
    FROM bulan b
    JOIN {tarikan} x
      ON x.ditarik_pada < ((b.awal_bulan + INTERVAL '1 month')::timestamp
                           AT TIME ZONE %(zona)s)
    GROUP BY b.awal_bulan, x.kantor_id"""


# Satu baris per (bulan, kantor). Agregat dihitung sekali per tarikan, baru
# dipilih per bulan — view hist_* tidak dipindai ulang untuk tiap bulan.
SQL = f"""
WITH
  capil_tarikan AS (
    SELECT kantor_id, tarikan_id, MIN(ditarik_pada) AS ditarik_pada,
      COUNT(*) FILTER (WHERE trim("K") = '1') AS aktif_1,
      COUNT(*) FILTER (WHERE trim("K") = '2') AS aktif_2,
      COUNT(*) FILTER (WHERE trim("K") = '3') AS aktif_3,
      COUNT(*) FILTER (WHERE "LMG" NOT LIKE 'PRA' AND "LMG" NOT LIKE 'PJ%%'
                         AND "LMG" NOT LIKE 'KPJ%%') AS petugas
    FROM {rs._q(rs.VIEW_CAPIL)}
    GROUP BY kantor_id, tarikan_id
  ),
  rekap_tarikan AS (
    SELECT kantor_id, tarikan_id, MIN(ditarik_pada) AS ditarik_pada,
      SUM(total_100_persen) FILTER (WHERE j = 'NOMINAL IFQ') AS wajib,
      SUM(total_100_persen) FILTER (WHERE j = 'NOMINAL ZKT') AS pokok,
      SUM(total_100_persen) FILTER (WHERE j = 'NOMINAL SDQ') AS sukarela,
      SUM(disetor)          FILTER (WHERE j <> 'TOTAL')      AS setor,
      SUM(disetor)          FILTER (WHERE j = 'CAD')         AS cad
    FROM (SELECT *, upper(trim(jenis)) AS j FROM {rs._q(rs.VIEW_REKAP)}) r
    GROUP BY kantor_id, tarikan_id
  ),
  rincian_tarikan AS (
    SELECT kantor_id, tarikan_id, MIN(ditarik_pada) AS ditarik_pada
    FROM {rs._q(rs.VIEW_RINCIAN)}
    GROUP BY kantor_id, tarikan_id
  ),
  bulan AS (
    SELECT generate_series(
      LEAST(%(awal)s::timestamp,
            (SELECT date_trunc('month', MIN(ditarik_pada) AT TIME ZONE %(zona)s)
             FROM rekap_tarikan)),
      %(akhir)s::timestamp, INTERVAL '1 month')::date AS awal_bulan
  ),
  p_capil   AS ({_pilih("capil_tarikan")}
  ),
  p_rekap   AS ({_pilih("rekap_tarikan")}
  ),
  p_rincian AS ({_pilih("rincian_tarikan")}
  ),
  tunai AS (
    SELECT pr.awal_bulan, pr.kantor_id, SUM(h.tunai_ifq) AS tunai
    FROM p_rincian pr
    JOIN {rs._q(rs.VIEW_RINCIAN)} h
      ON h.kantor_id = pr.kantor_id AND h.tarikan_id = pr.tarikan_id
    GROUP BY pr.awal_bulan, pr.kantor_id
  ),
  kunci AS (
    SELECT awal_bulan, kantor_id FROM p_capil
    UNION SELECT awal_bulan, kantor_id FROM p_rekap
    UNION SELECT awal_bulan, kantor_id FROM p_rincian
  )
SELECT
  k.awal_bulan, k.kantor_id,
  t.tunai,
  c.aktif_1, c.aktif_2, c.aktif_3, c.petugas,
  r.wajib, r.pokok, r.sukarela, r.setor, r.cad,
  c.ditarik_pada  AT TIME ZONE %(zona)s AS ditarik_capil,
  r.ditarik_pada  AT TIME ZONE %(zona)s AS ditarik_finance
FROM kunci k
  LEFT JOIN p_capil pc       ON pc.awal_bulan = k.awal_bulan AND pc.kantor_id = k.kantor_id
  LEFT JOIN capil_tarikan c  ON c.kantor_id = pc.kantor_id AND c.tarikan_id = pc.tarikan_id
  LEFT JOIN p_rekap pk       ON pk.awal_bulan = k.awal_bulan AND pk.kantor_id = k.kantor_id
  LEFT JOIN rekap_tarikan r  ON r.kantor_id = pk.kantor_id AND r.tarikan_id = pk.tarikan_id
  LEFT JOIN tunai t          ON t.awal_bulan = k.awal_bulan AND t.kantor_id = k.kantor_id
ORDER BY k.awal_bulan, k.kantor_id
"""

# Judul header template → nilai yang mengisi kolom di bawahnya, kiri ke kanan.
# Judul yang di-merge melebar (NASABAH AKTIF = 3 kolom) harus selebar daftarnya.
# Turunan (total_anggota, kontribusi, akumulasi_cad) dihitung di _nilai_baris().
HEADER = {
    "NASABAH TUNAI":     ["tunai"],
    "NASABAH AKTIF":     ["aktif_1", "aktif_2", "aktif_3"],
    "TOTAL ANGGOTA":     ["total_anggota"],
    "SIMPANAN WAJIB":    ["wajib"],
    "SIMPANAN POKOK":    ["pokok"],
    "SIMPANAN SUKARELA": ["sukarela"],
    "TOTAL KONTRIBUSI":  ["kontribusi"],
    "TOTAL SETOR":       ["setor"],
    "CAD BULAN INI":     ["cad"],
    "AKUMULASI CAD":     ["akumulasi_cad"],
    "JUMLAH PETUGAS":    ["petugas"],
}
NILAI = [n for daftar in HEADER.values() for n in daftar]
UANG = {"wajib", "pokok", "sukarela", "kontribusi", "setor", "cad", "akumulasi_cad"}

KOL_JUDUL_BULAN = "A2"    # baris kedua judul (merge selebar tabel), kosong di template
BARIS_HEADER    = 4
BARIS_PERTAMA   = 6       # baris data pertama; baris 1-5 judul + header dua tingkat
LABEL_TOTAL     = "TOTAL"
LABEL_BULAN     = "BULAN"


def laporan_aktif() -> tuple:
    bisa, alasan = rs.laporan_aktif()
    if not bisa:
        return bisa, alasan
    if not TEMPLATE.is_file():
        return False, f"template tidak ditemukan: {TEMPLATE}"
    return True, ""


def _uang(nilai) -> float:
    # DISETOR membawa noise float dari Excel (51677862.50000001).
    return round(float(nilai or 0), 2)


def ambil_data(bulan: int, tahun: int) -> dict:
    """{(tahun, bulan): {kantor_id: rec}} untuk semua bulan s.d. bulan laporan."""
    try:
        with rs.psycopg2.connect(connect_timeout=10, **rs.DB) as conn, conn.cursor() as cur:
            cur.execute(SQL, {
                "awal": date(tahun, 1, 1).isoformat(),
                "akhir": date(tahun, bulan, 1).isoformat(),
                "zona": rs.ZONA_WAKTU,
            })
            nama_kolom = [d[0] for d in cur.description]
            baris = cur.fetchall()
    except rs.psycopg2.Error as exc:
        raise ReportError(f"query warehouse gagal: {str(exc).strip()[:300]}") from exc

    hasil = {}
    for r in baris:
        rec = dict(zip(nama_kolom, r))
        awal = rec.pop("awal_bulan")
        kantor = str(rec.pop("kantor_id") or "").strip().upper()
        if kantor:
            hasil.setdefault((awal.year, awal.month), {})[kantor] = rec
    return hasil


def _nilai_baris(rec: dict, akumulasi: float) -> dict:
    """Nilai satu baris (kunci = nama di HEADER), termasuk kolom turunan."""
    n = {k: rs._angka(rec.get(k)) for k in
         ("tunai", "aktif_1", "aktif_2", "aktif_3", "petugas")}
    for k in ("wajib", "pokok", "sukarela", "setor", "cad"):
        n[k] = _uang(rec.get(k))
    n["total_anggota"] = n["aktif_1"] + n["aktif_2"] + n["aktif_3"]
    n["kontribusi"] = round(n["wajib"] + n["pokok"] + n["sukarela"], 2)
    n["akumulasi_cad"] = round(akumulasi, 2)
    return n


def _cari_baris(ws, label: str, mulai: int) -> int:
    for baris in range(mulai, ws.max_row + 1):
        if str(ws[f"A{baris}"].value or "").strip().upper() == label:
            return baris
    raise ReportError(f"template tidak punya baris berlabel {label!r} di kolom A "
                      f"— periksa {TEMPLATE}")


def _peta_kolom(ws) -> dict:
    """{nama nilai: huruf kolom}, dari judul di baris BARIS_HEADER."""
    lebar = {}      # kolom kiri sebuah merge → jumlah kolomnya
    for m in ws.merged_cells.ranges:
        if m.min_row <= BARIS_HEADER <= m.max_row:
            lebar[m.min_col] = m.max_col - m.min_col + 1
    peta, ketemu = {}, set()
    for kol in range(2, ws.max_column + 1):
        judul = " ".join(str(ws.cell(BARIS_HEADER, kol).value or "").split()).upper()
        if judul not in HEADER:
            continue
        daftar = HEADER[judul]
        if lebar.get(kol, 1) != len(daftar):
            raise ReportError(f"header {judul!r} di template selebar {lebar.get(kol, 1)} "
                              f"kolom, seharusnya {len(daftar)} — periksa {TEMPLATE}")
        for i, nama in enumerate(daftar):
            peta[nama] = get_column_letter(kol + i)
        ketemu.add(judul)
    hilang = [j for j in HEADER if j not in ketemu]
    if hilang:
        raise ReportError("header tidak ditemukan di baris "
                          f"{BARIS_HEADER} template: {', '.join(hilang)} — periksa {TEMPLATE}")
    return peta


def _tulis(ws, kolom: dict, baris: int, nilai: dict) -> None:
    # Total ditulis sebagai angka, menggantikan =SUM(...) template: pratinjau
    # file Telegram tidak menghitung ulang formula.
    for nama, kol in kolom.items():
        v = nilai.get(nama, 0)
        ws[f"{kol}{baris}"] = round(v, 2) if nama in UANG else v


def bangun_laporan(bulan: int, tahun: int) -> tuple:
    """Kembalikan (isi_xlsx_bytes, nama_file, catatan) — kontrak report_summary."""
    bisa, alasan = laporan_aktif()
    if not bisa:
        raise ReportError(alasan)
    bulan, tahun = int(bulan), int(tahun)
    if not 1 <= bulan <= 12:
        raise ReportError(f"bulan harus 1-12, bukan {bulan}")

    data = ambil_data(bulan, tahun)

    wb = rs.load_workbook(TEMPLATE)
    ws = wb.active
    ws.title = f"{bulan}_{tahun % 100:02d}"
    ws[KOL_JUDUL_BULAN] = f"BULAN: {NAMA_BULAN[bulan]} {tahun}"

    kolom = _peta_kolom(ws)
    baris_total = _cari_baris(ws, LABEL_TOTAL, BARIS_PERTAMA)
    kantor_template = {}
    for baris in range(BARIS_PERTAMA, baris_total):
        kode = str(ws[f"A{baris}"].value or "").strip().upper()
        if kode:
            kantor_template[kode] = baris

    # Semua bulan berurutan: akumulasi CAD per kantor berjalan dari bulan
    # tarikan pertama, menyeberang tahun bila perlu.
    akumulasi = {k: 0.0 for k in kantor_template}
    per_bulan = {}      # bulan di tahun laporan → {kantor: nilai_baris}
    for (t, b) in sorted(data):
        isi = {}
        for kantor in kantor_template:
            rec = data[(t, b)].get(kantor)
            if rec is None:
                continue
            # Hanya tarikan finance dari bulan itu sendiri yang menambah
            # akumulasi. Kantor tanpa tarikan baru memakai potret bulan lalu
            # (seperti kolom lain), tapi CAD-nya bukan uang baru — kalau
            # ditambahkan lagi, akumulasinya terhitung dua kali.
            waktu = rec.get("ditarik_finance")
            if waktu is not None and (waktu.year, waktu.month) == (t, b):
                akumulasi[kantor] += _uang(rec.get("cad"))
            isi[kantor] = _nilai_baris(rec, akumulasi[kantor])
        if t == tahun:
            per_bulan[b] = isi

    # Tabel atas: bulan laporan per kantor.
    ini = per_bulan.get(bulan, {})
    total = {nama: 0 for nama in NILAI}
    for kantor, baris in kantor_template.items():
        nilai = ini.get(kantor)
        if nilai is None:
            continue    # kantor tanpa data: barisnya dibiarkan kosong
        _tulis(ws, kolom, baris, nilai)
        for nama in NILAI:
            total[nama] += nilai[nama]
    _tulis(ws, kolom, baris_total, total)

    # Tabel bawah: total semua kantor per bulan di tahun laporan.
    baris_judul = _cari_baris(ws, LABEL_BULAN, baris_total + 1)
    baris_total_bawah = _cari_baris(ws, LABEL_TOTAL, baris_judul + 1)
    total_bawah = {nama: 0 for nama in NILAI}
    for baris in range(baris_judul + 1, baris_total_bawah):
        label = str(ws[f"A{baris}"].value or "").strip().upper()
        if label not in NAMA_BULAN or not per_bulan.get(NAMA_BULAN.index(label)):
            continue    # bulan sesudah bulan laporan / tanpa data: kosong
        sebulan = {nama: sum(v[nama] for v in per_bulan[NAMA_BULAN.index(label)].values())
                   for nama in NILAI}
        _tulis(ws, kolom, baris, sebulan)
        for nama, v in sebulan.items():
            total_bawah[nama] += v
    _tulis(ws, kolom, baris_total_bawah, total_bawah)

    _tambah_grafik(wb, ws, kolom, bulan, tahun,
                   kantor=(BARIS_PERTAMA, baris_total - 1),
                   bulanan=(baris_judul + 1, baris_total_bawah - 1))

    catatan = _catatan(data.get((tahun, bulan), {}), kantor_template, bulan, tahun)

    buf = io.BytesIO()
    wb.save(buf)
    nama_file = f"keuangan-{tahun:04d}-{bulan:02d}.xlsx"
    log.info("Laporan %s dibangun (%d kantor terisi, %d catatan).",
             nama_file, len(ini), len(catatan))
    return rs._kunci_xlsx(buf.getvalue()), nama_file, catatan


# ── Grafik ──────────────────────────────────────────────────────────────────

# Warna kategori dipakai berurutan, tidak pernah diputar ulang: seri pertama
# selalu biru, kedua oranye, ketiga hijau-tosca (K1/K2/K3, wajib/pokok/sukarela).
WARNA = ["2A78D6", "EB6834", "1BAF7A"]
WARNA_GRID = "E1E0DC"
WARNA_LATAR = "FFFFFF"
FMT_RUPIAH = '#,##0.0,," jt"'     # 93.959.750 → "94,0 jt"
FMT_ORANG = "#,##0"

# Tata letak sheet GRAFIK: dua kolom grafik, lebar kolom sel dikunci supaya
# jarak antargrafik tidak bergantung pada lebar kolom bawaan aplikasi.
LEBAR_CM, TINGGI_CM, TINGGI_KANTOR_CM = 16, 7.5, 11
LEBAR_KOLOM_SEL = 9.3             # ≈ 70 px; 9 kolom ≈ 16,5 cm
JANGKAR_KOLOM = ["A", "J"]
BARIS_PER_TINGKAT = 16            # baris 15 pt → 16 baris ≈ 8,5 cm


def _gaya_sumbu(chart, fmt: str) -> None:
    chart.x_axis.delete = False     # openpyxl 3.1 menyembunyikan sumbu bila tidak diisi
    chart.y_axis.delete = False
    chart.y_axis.numFmt = fmt
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(
        ln=LineProperties(solidFill=WARNA_GRID, w=6350))
    chart.y_axis.scaling.min = 0
    chart.width, chart.height = LEBAR_CM, TINGGI_CM


def _ref(ws, kolom: str, rentang: tuple) -> "Reference":
    # Bentuk koordinat, bukan range_string: nama sheet "10_26" perlu dikutip.
    i = column_index_from_string(kolom)
    return Reference(ws, min_col=i, max_col=i, min_row=rentang[0], max_row=rentang[1])


def _kolom(ws, judul, seri, kategori, rentang, fmt, *, tumpuk=False, mendatar=False):
    """Diagram batang; `seri` = [(huruf kolom, nama)]. Legend hanya bila >1 seri."""
    c = BarChart()
    c.type = "bar" if mendatar else "col"
    c.title = judul
    if tumpuk:
        c.grouping, c.overlap = "stacked", 100
    c.gapWidth = 60
    for i, (kol, nama) in enumerate(seri):
        s = Series(_ref(ws, kol, rentang), title=nama)
        s.graphicalProperties.solidFill = WARNA[i]
        # Garis tepi warna latar = celah tipis antar segmen / batang.
        s.graphicalProperties.line.solidFill = WARNA_LATAR
        s.graphicalProperties.line.width = 12700
        c.series.append(s)
    c.set_categories(kategori)
    _gaya_sumbu(c, fmt)
    if mendatar:
        c.x_axis.scaling.orientation = "maxMin"   # kantor pertama di atas
        c.x_axis.tickLblSkip = 1                  # semua kode kantor tampil
        c.height = TINGGI_KANTOR_CM
    if len(seri) > 1:
        c.legend.position = "b"
    else:
        c.legend = None             # satu seri: judul sudah menamainya
    return c


def _garis(ws, judul, kolom, kategori, rentang, fmt):
    """Satu seri garis per grafik — skala wajib/pokok/sukarela terlalu berbeda
    untuk berbagi sumbu (sukarela akan rata di dasar)."""
    c = LineChart()
    c.title = judul
    s = Series(_ref(ws, kolom, rentang), title=judul)
    s.graphicalProperties.line.solidFill = WARNA[0]
    s.graphicalProperties.line.width = 25400
    s.smooth = False
    s.marker.symbol, s.marker.size = "circle", 7
    s.marker.graphicalProperties = GraphicalProperties(solidFill=WARNA[0])
    s.marker.graphicalProperties.line.solidFill = WARNA_LATAR
    c.series.append(s)
    c.set_categories(kategori)
    c.display_blanks = "gap"        # bulan sesudah bulan laporan = kosong, bukan nol
    _gaya_sumbu(c, fmt)
    c.legend = None
    return c


def _tambah_grafik(wb, ws, kolom: dict, bulan: int, tahun: int, *,
                   kantor: tuple, bulanan: tuple) -> None:
    """
    Sheet GRAFIK: tren bulanan (tabel bawah) dan perbandingan kantor bulan
    laporan (tabel atas). Kolom dan baris yang dirujuk sama dengan yang diisi —
    keduanya dicari dari label template.
    """
    g = wb.create_sheet("GRAFIK")
    g["A1"] = f"GRAFIK REKAPITULASI SIMPANAN — {NAMA_BULAN[bulan]} {tahun}"
    g["A1"].font = Font(bold=True, size=14)
    for i in range(1, 19):
        g.column_dimensions[get_column_letter(i)].width = LEBAR_KOLOM_SEL

    bln = _ref(ws, "A", bulanan)
    ktr = _ref(ws, "A", kantor)
    aktif = [(kolom[f"aktif_{i}"], f"K{i}") for i in (1, 2, 3)]
    label = f"{NAMA_BULAN[bulan]} {tahun}"

    grafik = [
        _garis(ws, "Nasabah Tunai per Bulan", kolom["tunai"], bln, bulanan, FMT_ORANG),
        _kolom(ws, "Nasabah Aktif per Bulan", aktif, bln, bulanan, FMT_ORANG, tumpuk=True),
        _garis(ws, "Simpanan Wajib per Bulan", kolom["wajib"], bln, bulanan, FMT_RUPIAH),
        _garis(ws, "Simpanan Pokok per Bulan", kolom["pokok"], bln, bulanan, FMT_RUPIAH),
        _garis(ws, "Simpanan Sukarela per Bulan", kolom["sukarela"], bln, bulanan, FMT_RUPIAH),
        _kolom(ws, "Total Kontribusi vs Total Setor per Bulan",
               [(kolom["kontribusi"], "Total kontribusi"), (kolom["setor"], "Total setor")],
               bln, bulanan, FMT_RUPIAH),
        _kolom(ws, "CAD Bulan Ini per Bulan", [(kolom["cad"], "CAD bulan ini")],
               bln, bulanan, FMT_RUPIAH),
        _garis(ws, "Akumulasi CAD", kolom["akumulasi_cad"], bln, bulanan, FMT_RUPIAH),
        _kolom(ws, f"Komposisi Simpanan per Kantor — {label}",
               [(kolom["wajib"], "Wajib"), (kolom["pokok"], "Pokok"),
                (kolom["sukarela"], "Sukarela")],
               ktr, kantor, FMT_RUPIAH, tumpuk=True, mendatar=True),
        _kolom(ws, f"Nasabah Tunai vs Total Anggota per Kantor — {label}",
               [(kolom["tunai"], "Nasabah tunai"), (kolom["total_anggota"], "Total anggota")],
               ktr, kantor, FMT_ORANG, mendatar=True),
    ]
    for i, c in enumerate(grafik):
        g.add_chart(c, f"{JANGKAR_KOLOM[i % 2]}{3 + (i // 2) * BARIS_PER_TINGKAT}")

    # Cetak: A4 tegak, dua grafik selebar satu halaman; tanpa ini grafik
    # terpotong di tengah oleh batas halaman bawaan.
    g.page_setup.paperSize = g.PAPERSIZE_A4
    g.page_setup.orientation = "portrait"
    g.sheet_properties.pageSetUpPr.fitToPage = True
    g.page_setup.fitToWidth, g.page_setup.fitToHeight = 1, 0


def _catatan(data_bulan: dict, kantor_template: dict, bulan: int, tahun: int) -> list:
    """Hal yang perlu diketahui operator; ditulis ke log relay, bukan ke Telegram."""
    catatan = rs._catatan_tarikan(data_bulan, bulan, tahun)
    kosong = sorted(set(kantor_template) - set(data_bulan))
    if kosong:
        catatan.append("Tanpa data di warehouse: " + ", ".join(kosong))
    lebih = sorted(set(data_bulan) - set(kantor_template))
    if lebih:
        catatan.append("Ada di warehouse tapi tidak ada barisnya di template: "
                       + ", ".join(lebih))
    return catatan
