#!/usr/bin/env python3
"""
report_summary.py — Laporan Rekap Bulanan (satu file Excel)

Membaca angka rekap per kantor dari warehouse (schema `analytics`) lalu
menuangkannya ke template `docs/template/summary.xlsx`. Templatenya yang
dipakai apa adanya — merge, border, dan format angkanya sudah sesuai bentuk
laporan yang dipakai manual selama ini, jadi modul ini hanya *mengisi sel*,
tidak pernah menggambar ulang tabelnya.

Dipakai oleh notif_relay.py: perintah /laporan dan tombol Laporan di Mini App
memanggil bangun_laporan(), lalu hasilnya dikirim ke Telegram sebagai dokumen.

Yang perlu diingat soal cakupan datanya
---------------------------------------
Model staging hanya menyimpan *tarikan terakhir* (`_airbyte_generation_id`
= MAX), jadi warehouse tidak menyimpan riwayat bulanan. Konsekuensinya:

  * Semua kolom selain DAKWAH HASIL adalah potret data terkini, bukan potret
    bulan yang dipilih. Bulan/tahun yang dipilih hanya menyaring DAKWAH HASIL
    (Bln_Integrasi/Th_Integrasi) dan menjadi judul laporan.
  * Baris "JUMLAH BULAN LALU" dan "SELISIH" diambil dari *arsip* laporan bulan
    sebelumnya (REPORT_HISTORY_DIR), bukan dari warehouse. Tiap laporan yang
    selesai dibangun menyimpan potretnya sendiri, jadi pembanding baru ada
    setelah bulan itu pernah dilaporkan sekali. Bulan pertama akan berisi nol,
    dan itu disebut di catatan laporannya.

Dua hal kosmetik hilang saat openpyxl menulis ulang template: data bar
conditional formatting (semuanya min=0/max=0, jadi tidak menggambar apa pun) dan
sisa threaded comment lama. Struktur tabel — merge, border, lebar kolom, format
angka — tetap utuh.

Konfigurasi via environment:
  REPORT_DB_HOST      (default postgres-db)
  REPORT_DB_PORT      (default 5432)
  REPORT_DB_NAME      (default capil_db)
  REPORT_DB_USER      (default admin)
  REPORT_DB_PASSWORD  (default password123)
  REPORT_SCHEMA       (default analytics)
  REPORT_VIEW_CAPIL           (default stg_all_capil)
  REPORT_VIEW_FINANCE_REKAP   (default stg_all_finance_rekap)
  REPORT_VIEW_FINANCE_RINCIAN (default stg_all_finance_rincian)
  REPORT_TEMPLATE     (default docs/template/summary.xlsx di sebelah modul ini)
  REPORT_HISTORY_DIR  (default /data/laporan) -- arsip potret bulanan
"""

import io
import json
import logging
import os
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

log = logging.getLogger("notif_relay.report")

# Impor berat dibuat opsional supaya relay tetap hidup di image lama yang belum
# punya psycopg2/openpyxl — fiturnya mati, sisanya jalan seperti biasa.
try:
    import psycopg2
except ImportError:  # pragma: no cover
    psycopg2 = None
try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None


# ── Konfigurasi ─────────────────────────────────────────────────────────────

DB = {
    "host":     os.getenv("REPORT_DB_HOST", "postgres-db"),
    "port":     int(os.getenv("REPORT_DB_PORT", "5432")),
    "dbname":   os.getenv("REPORT_DB_NAME", "capil_db"),
    "user":     os.getenv("REPORT_DB_USER", "admin"),
    "password": os.getenv("REPORT_DB_PASSWORD", "password123"),
}

SCHEMA          = os.getenv("REPORT_SCHEMA", "analytics")
VIEW_CAPIL      = os.getenv("REPORT_VIEW_CAPIL", "stg_all_capil")
VIEW_REKAP      = os.getenv("REPORT_VIEW_FINANCE_REKAP", "stg_all_finance_rekap")
VIEW_RINCIAN    = os.getenv("REPORT_VIEW_FINANCE_RINCIAN", "stg_all_finance_rincian")

TEMPLATE = Path(os.getenv(
    "REPORT_TEMPLATE",
    Path(__file__).resolve().parent.parent / "docs" / "template" / "summary.xlsx",
))

BULAN_SINGKAT = ["", "Jan", "Feb", "Mar", "Apr", "Mei", "Jun",
                 "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]

# Jenis di stg_all_finance_rekap yang masuk hitungan Infaq (INF ++).
JENIS_IP = "('NOMINAL AQQ','NOMINAL FDY','NOMINAL FI','NOMINAL IFQ'," \
           "'NOMINAL LQT','NOMINAL SDQ','NOMINAL SNK')"


def _q(nama: str) -> str:
    return f'"{SCHEMA}"."{nama}"'


# Query rekap. Sama persis dengan pertanyaan yang dipakai di Metabase, kecuali
# CTE `rekrut` yang bulan/tahunnya dijadikan parameter (%(bulan)s / %(tahun)s)
# alih-alih CURRENT_DATE - 1 bulan, supaya laporan bulan mana pun bisa diminta.
SQL = f"""
WITH
  pengurus AS (
    SELECT kantor_id,
      COUNT(*) FILTER (WHERE upper("JK") = 'L') AS pengurus_r,
      COUNT(*) FILTER (WHERE upper("JK") = 'P') AS pengurus_n,
      COUNT(*) AS pengurus_jumlah,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'A')  AS pengurus_aktivitas_a,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'M')  AS pengurus_aktivitas_m,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'AM') AS pengurus_aktivitas_am,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'NA') AS pengurus_aktivitas_na
    FROM {_q(VIEW_CAPIL)}
    WHERE "LMG" NOT LIKE 'PRA' AND "LMG" NOT LIKE 'PJ%%' AND "LMG" NOT LIKE 'KPJ%%'
    GROUP BY kantor_id
  ),
  anggota AS (
    SELECT kantor_id,
      COUNT(*) FILTER (WHERE upper("JK") = 'L') AS anggota_r,
      COUNT(*) FILTER (WHERE upper("JK") = 'P') AS anggota_n,
      COUNT(*) AS anggota_jumlah,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'A')  AS anggota_aktivitas_a,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'M')  AS anggota_aktivitas_m,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'AM') AS anggota_aktivitas_am,
      COUNT(*) FILTER (WHERE "Status_Aktivitas" = 'NA') AS anggota_aktivitas_na
    FROM {_q(VIEW_CAPIL)}
    WHERE "LMG" LIKE 'PRA' OR "LMG" LIKE 'PJ%%' OR "LMG" LIKE 'KPJ%%'
    GROUP BY kantor_id
  ),
  jenjang AS (
    SELECT kantor_id,
      COUNT(*) FILTER (WHERE "K" = '1') AS jenjang_a1,
      COUNT(*) FILTER (WHERE "K" = '2') AS jenjang_a2,
      COUNT(*) FILTER (WHERE "K" = '3') AS jenjang_a3
    FROM {_q(VIEW_CAPIL)}
    GROUP BY kantor_id
  ),
  rekrut AS (
    SELECT kantor_id, COUNT(*) AS rekrut
    FROM {_q(VIEW_CAPIL)}
    WHERE "Bln_Integrasi"::int = %(bulan)s
      AND "Th_Integrasi"::int  = %(tahun)s
    GROUP BY kantor_id
  ),
  total_tunai AS (
    SELECT kantor_id,
      SUM(COALESCE("tunai_fi",0)) + SUM(COALESCE("tunai_aqq",0))
      + SUM(COALESCE("tunai_ifq",0)) + SUM(COALESCE("tunai_lqt",0))
      + SUM(COALESCE("tunai_sdq",0)) + SUM(COALESCE("tunai_snk",0))
      + SUM(COALESCE("tunai_fdy",0)) AS mfq
    FROM {_q(VIEW_RINCIAN)}
    GROUP BY kantor_id
  ),
  income_ip AS (
    SELECT kantor_id,
      SUM(CASE WHEN jenis IN {JENIS_IP} THEN total_100_persen ELSE 0 END)
        AS ip_penerimaan_100_persen
    FROM {_q(VIEW_REKAP)}
    GROUP BY kantor_id
  ),
  transfer_ip AS (
    SELECT kantor_id,
      SUM(CASE WHEN jenis IN {JENIS_IP} THEN pembulatan_setor ELSE 0 END)
        AS ip_terima_dpp
    FROM {_q(VIEW_REKAP)}
    GROUP BY kantor_id
  ),
  income_sd AS (
    SELECT kantor_id, SUM(COALESCE(nominal_zkt,0)) AS sd_penerimaan_100_persen
    FROM {_q(VIEW_RINCIAN)}
    GROUP BY kantor_id
  ),
  transfer_sd AS (
    SELECT kantor_id,
      SUM(CASE WHEN jenis IN ('NOMINAL ZKT') THEN pembulatan_setor ELSE 0 END)
        AS sd_terima_dpp
    FROM {_q(VIEW_REKAP)}
    GROUP BY kantor_id
  ),
  cadangan AS (
    SELECT kantor_id,
      SUM(CASE WHEN jenis IN ('CAD ') THEN pembulatan_setor ELSE 0 END) AS cad
    FROM {_q(VIEW_REKAP)}
    GROUP BY kantor_id
  ),
  tunai_lainnya AS (
    SELECT kantor_id,
      SUM(COALESCE("tunai_zf",0)) + SUM(COALESCE("tunai_tdy",0))
        AS lainnya_tunai_jiwa
    FROM {_q(VIEW_RINCIAN)}
    GROUP BY kantor_id
  ),
  income_lainnya AS (
    SELECT kantor_id,
      SUM(COALESCE(nominal_zf,0) + COALESCE(nominal_tdy,0))
        AS lainnya_penerimaan_100_persen
    FROM {_q(VIEW_RINCIAN)}
    GROUP BY kantor_id
  ),
  transfer_lainnya AS (
    SELECT kantor_id,
      SUM(CASE WHEN jenis IN ('NOMINAL ZF','NOMINAL TDY') THEN pembulatan_setor
               ELSE 0 END) AS lainnya_terima_dpp
    FROM {_q(VIEW_REKAP)}
    GROUP BY kantor_id
  )
SELECT
  g.kantor_id,
  g.pengurus_r, g.pengurus_n, g.pengurus_jumlah,
  g.pengurus_aktivitas_a, g.pengurus_aktivitas_m,
  g.pengurus_aktivitas_am, g.pengurus_aktivitas_na,
  a.anggota_r, a.anggota_n, a.anggota_jumlah,
  a.anggota_aktivitas_a, a.anggota_aktivitas_m,
  a.anggota_aktivitas_am, a.anggota_aktivitas_na,
  j.jenjang_a1, j.jenjang_a2, j.jenjang_a3,
  r.rekrut,
  t.mfq,
  i.ip_penerimaan_100_persen, tfip.ip_terima_dpp,
  isd.sd_penerimaan_100_persen, tfsd.sd_terima_dpp,
  c.cad,
  tl.lainnya_tunai_jiwa, il.lainnya_penerimaan_100_persen, tfl.lainnya_terima_dpp
FROM pengurus g
  LEFT JOIN anggota          a    ON a.kantor_id    = g.kantor_id
  LEFT JOIN jenjang          j    ON j.kantor_id    = g.kantor_id
  LEFT JOIN rekrut           r    ON r.kantor_id    = g.kantor_id
  LEFT JOIN total_tunai      t    ON t.kantor_id    = g.kantor_id
  LEFT JOIN income_ip        i    ON i.kantor_id    = g.kantor_id
  LEFT JOIN transfer_ip      tfip ON tfip.kantor_id = g.kantor_id
  LEFT JOIN income_sd        isd  ON isd.kantor_id  = g.kantor_id
  LEFT JOIN transfer_sd      tfsd ON tfsd.kantor_id = g.kantor_id
  LEFT JOIN cadangan         c    ON c.kantor_id    = g.kantor_id
  LEFT JOIN tunai_lainnya    tl   ON tl.kantor_id   = g.kantor_id
  LEFT JOIN income_lainnya   il   ON il.kantor_id   = g.kantor_id
  LEFT JOIN transfer_lainnya tfl  ON tfl.kantor_id  = g.kantor_id
"""

# Kolom template → nama kolom hasil query. Urutan header template (baris 2-4)
# yang menentukan, bukan urutan SELECT.
#
# Sengaja TIDAK diisi:
#   D  "BARIS"                — tidak ada di query, biarkan seperti template
#   S  "NT" (aktivitas anggota) — idem
KOLOM = {
    "E":  "pengurus_r",                  # PENGURUS · R
    "F":  "pengurus_n",                  # PENGURUS · N
    "G":  "pengurus_jumlah",             # PENGURUS · JLH
    "H":  "pengurus_aktivitas_a",        # AKTIVITAS PENGURUS · A
    "I":  "pengurus_aktivitas_m",        # · M
    "J":  "pengurus_aktivitas_am",       # · AM
    "K":  "pengurus_aktivitas_na",       # · NA
    "L":  "anggota_r",                   # ANGGOTA · R
    "M":  "anggota_n",                   # ANGGOTA · N
    "N":  "anggota_jumlah",              # ANGGOTA · JLH
    "O":  "anggota_aktivitas_a",         # AKTIVITAS ANGGOTA · A
    "P":  "anggota_aktivitas_m",         # · M
    "Q":  "anggota_aktivitas_am",        # · AM
    "R":  "anggota_aktivitas_na",        # · NA
    "U":  "jenjang_a1",                  # JENJANG · A1
    "V":  "jenjang_a2",                  # JENJANG · A2
    "W":  "jenjang_a3",                  # JENJANG · A3
    "X":  "rekrut",                      # DAKWAH · HASIL
    "Y":  "mfq",                         # MFQ
    "Z":  "ip_penerimaan_100_persen",    # INF ++ · PENERIMAAN 100%
    "AA": "ip_terima_dpp",               # INF ++ · TERIMA DPP
    "AB": "sd_penerimaan_100_persen",    # SD · PENERIMAAN 100%
    "AC": "sd_terima_dpp",               # SD · TERIMA DPP
    "AD": "cad",                         # CAD
    "AE": "lainnya_tunai_jiwa",          # ZF · TUNAI JIWA
    "AF": "lainnya_penerimaan_100_persen",  # ZF · PENERIMAAN 100%
    "AG": "lainnya_terima_dpp",          # ZF · TERIMA DPP
}

# T = P+A: jumlah pengurus + anggota, tidak ada di query (di Metabase dibaca
# dari dua kolom terpisah).
KOLOM_TURUNAN = {"T": ("pengurus_jumlah", "anggota_jumlah")}

# Kolom yang totalnya ditulis di baris JUMLAH BULAN INI. D dan S ikut meski tidak
# pernah diisi: templatenya menaruh =SUM(...) di seluruh baris itu, dan formula
# yang tersisa akan tampil sebagai hasil lama (0) di previewer yang tidak
# menghitung ulang — termasuk pratinjau file bawaan Telegram.
KOLOM_TOTAL = ["D", "S"] + list(KOLOM) + list(KOLOM_TURUNAN)

# Nama semantik tiap kolom, dipakai sebagai kunci arsip bulanan. Sengaja BUKAN
# huruf kolom: kalau suatu saat template menyisipkan kolom baru, huruf-hurufnya
# bergeser dan arsip lama akan dibaca ke kolom yang salah tanpa error.
NAMA_KOLOM = dict(KOLOM)
NAMA_KOLOM.update({"T": "p_plus_a", "D": "baris", "S": "nt"})

KOL_JUDUL     = "A1"    # "BULAN: Apr 26"
KOL_KANTOR    = "C"     # kode kantor per baris
BARIS_PERTAMA = 5       # baris data pertama; baris 1-4 adalah blok judul + header

# Baris ringkasan dicari dari label di kolom A, bukan dikunci ke nomor baris:
# menambah kantor ke template menggeser semuanya, dan nomor baris yang salah
# menulis angka ke tengah tabel tanpa ada yang mengeluh.
LABEL_TOTAL   = "JUMLAH BULAN INI"
LABEL_LALU    = "JUMLAH BULAN LALU"
LABEL_SELISIH = "SELISIH"

# Arsip bulanan. Tiap laporan yang selesai dibangun disimpan di sini, dan
# laporan bulan berikutnya membacanya untuk mengisi baris JUMLAH BULAN LALU.
# Warehouse tidak menyimpan riwayat, jadi arsip inilah satu-satunya sumber
# angka bulan lalu — taruh di volume yang ikut backup.
HISTORY_DIR = Path(os.getenv("REPORT_HISTORY_DIR", "/data/laporan"))


class ReportError(Exception):
    """Kegagalan yang layak ditampilkan apa adanya ke operator di Telegram."""


def laporan_aktif() -> tuple:
    """(bisa_dipakai, alasan). Alasan hanya berarti kalau bisa_dipakai False."""
    if psycopg2 is None or load_workbook is None:
        return False, ("image notif-relay ini belum punya psycopg2/openpyxl — "
                       "build ulang: docker compose build notif-relay")
    if not TEMPLATE.is_file():
        return False, f"template tidak ditemukan: {TEMPLATE}"
    return True, ""


def bulan_default() -> tuple:
    """Bulan laporan default: bulan lalu (bulan berjalan biasanya belum lengkap)."""
    hari_ini = date.today()
    return (12, hari_ini.year - 1) if hari_ini.month == 1 else (hari_ini.month - 1, hari_ini.year)


def _angka(nilai):
    """Rapikan nilai dari psycopg2 agar Excel menyimpannya sebagai angka."""
    if nilai is None:
        return 0
    if isinstance(nilai, Decimal):
        return int(nilai) if nilai == nilai.to_integral_value() else float(nilai)
    return nilai


def ambil_data(bulan: int, tahun: int) -> dict:
    """Jalankan query rekap; kembalikan {kantor_id: {kolom: nilai}}."""
    try:
        with psycopg2.connect(connect_timeout=10, **DB) as conn, conn.cursor() as cur:
            cur.execute(SQL, {"bulan": bulan, "tahun": tahun})
            nama_kolom = [d[0] for d in cur.description]
            baris = cur.fetchall()
    except psycopg2.Error as exc:
        # Pesan psycopg2 sudah menyebut relasi yang hilang / kredensial salah;
        # itu justru yang perlu dibaca operator di Telegram.
        raise ReportError(f"query warehouse gagal: {str(exc).strip()[:300]}") from exc

    hasil = {}
    for r in baris:
        rec = dict(zip(nama_kolom, r))
        kantor = str(rec.get("kantor_id") or "").strip().upper()
        if kantor:
            hasil[kantor] = {k: _angka(v) for k, v in rec.items() if k != "kantor_id"}
    return hasil


def _bulan_sebelum(bulan: int, tahun: int) -> tuple:
    return (12, tahun - 1) if bulan == 1 else (bulan - 1, tahun)


def _berkas_arsip(bulan: int, tahun: int) -> Path:
    return HISTORY_DIR / f"{tahun:04d}-{bulan:02d}.json"


def simpan_arsip(bulan: int, tahun: int, total: dict, per_kantor: dict) -> None:
    """
    Simpan potret satu bulan supaya laporan bulan berikutnya punya pembanding.

    Ditulis lewat file sementara lalu di-rename: laporan bisa dipicu berbarengan
    dari Telegram dan Mini App, dan arsip yang tertulis separuh akan merusak
    baris JUMLAH BULAN LALU bulan depan tanpa gejala apa pun.
    """
    HISTORY_DIR.mkdir(parents=True, exist_ok=True)
    isi = {
        "bulan": bulan,
        "tahun": tahun,
        "dibuat": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        # Kunci memakai nama semantik (NAMA_KOLOM), bukan huruf kolom.
        "total": {NAMA_KOLOM[k]: v for k, v in total.items() if k in NAMA_KOLOM},
        "kantor": per_kantor,
    }
    tujuan = _berkas_arsip(bulan, tahun)
    sementara = tujuan.with_suffix(".json.tmp")
    sementara.write_text(json.dumps(isi, indent=2), encoding="utf-8")
    sementara.replace(tujuan)
    log.info("Arsip %s disimpan.", tujuan)


def baca_arsip(bulan: int, tahun: int) -> dict:
    """Baca arsip satu bulan; {} kalau belum ada atau rusak."""
    berkas = _berkas_arsip(bulan, tahun)
    try:
        return json.loads(berkas.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}
    except (OSError, json.JSONDecodeError) as exc:
        log.warning("Arsip %s tidak terbaca: %s", berkas, exc)
        return {}


def _cari_baris_label(ws) -> dict:
    """
    Petakan label ringkasan di kolom A ke nomor barisnya.

    Dicari, bukan dikunci: template pernah bertambah dari 12 kantor ke 14, dan
    nomor baris yang tertinggal akan menimpa baris data dengan angka total.
    """
    incar = {LABEL_TOTAL: "total", LABEL_LALU: "lalu", LABEL_SELISIH: "selisih"}
    ketemu = {}
    for baris in range(BARIS_PERTAMA, ws.max_row + 1):
        teks = str(ws[f"A{baris}"].value or "").strip().upper()
        if teks in incar and incar[teks] not in ketemu:
            ketemu[incar[teks]] = baris
    if "total" not in ketemu:
        raise ReportError(
            f"template tidak punya baris berlabel {LABEL_TOTAL!r} di kolom A — "
            f"periksa {TEMPLATE}"
        )
    return ketemu


def bangun_laporan(bulan: int, tahun: int) -> tuple:
    """
    Isi template dengan data bulan/tahun tsb.

    Kembalikan (isi_xlsx_bytes, nama_file, catatan) — `catatan` adalah daftar
    hal yang perlu diketahui operator (kantor tanpa data, kantor yang ada di
    warehouse tapi tidak punya baris di template).
    """
    bisa, alasan = laporan_aktif()
    if not bisa:
        raise ReportError(alasan)
    if not 1 <= int(bulan) <= 12:
        raise ReportError(f"bulan harus 1-12, bukan {bulan}")
    bulan, tahun = int(bulan), int(tahun)

    data = ambil_data(bulan, tahun)

    wb = load_workbook(TEMPLATE)
    ws = wb.active
    ws.title = f"{bulan}_{tahun % 100:02d}"
    ws[KOL_JUDUL] = f"BULAN: {BULAN_SINGKAT[bulan]} {tahun % 100:02d}"

    baris_label = _cari_baris_label(ws)
    baris_terakhir = baris_label["total"] - 1

    total = {kol: 0 for kol in KOLOM_TOTAL}
    per_kantor, terpakai, kosong = {}, set(), []

    for baris in range(BARIS_PERTAMA, baris_terakhir + 1):
        kantor = str(ws[f"{KOL_KANTOR}{baris}"].value or "").strip().upper()
        if not kantor:
            continue
        rec = data.get(kantor)
        if rec is None:
            kosong.append(kantor)
            rec = {}
        else:
            terpakai.add(kantor)

        nilai_baris = {}
        for kol in ("D", "S"):   # tidak diisi query; totalnya tetap harus benar
            ada = ws[f"{kol}{baris}"].value
            nilai_baris[kol] = ada if isinstance(ada, (int, float)) else 0
        for kol, nama in KOLOM.items():
            nilai_baris[kol] = _angka(rec.get(nama))
        for kol, bagian in KOLOM_TURUNAN.items():
            nilai_baris[kol] = sum(_angka(rec.get(n)) for n in bagian)

        for kol, nilai in nilai_baris.items():
            if kol not in ("D", "S"):   # dua kolom itu dibiarkan apa adanya
                ws[f"{kol}{baris}"] = nilai
            total[kol] += nilai
        per_kantor[kantor] = {NAMA_KOLOM[k]: v for k, v in nilai_baris.items()}

    for kol, nilai in total.items():
        ws[f"{kol}{baris_label['total']}"] = nilai

    # Baris JUMLAH BULAN LALU + SELISIH, dari arsip bulan sebelumnya.
    b_lalu, t_lalu = _bulan_sebelum(bulan, tahun)
    arsip = baca_arsip(b_lalu, t_lalu)
    total_lalu = arsip.get("total") or {}
    catatan = []
    if not total_lalu:
        catatan.append(
            f"Belum ada arsip {BULAN_SINGKAT[b_lalu]} {t_lalu} — baris JUMLAH "
            f"BULAN LALU & SELISIH diisi 0."
        )
    for kunci, nomor in (("lalu", baris_label.get("lalu")),
                         ("selisih", baris_label.get("selisih"))):
        if not nomor:
            continue
        for kol, nilai_ini in total.items():
            # Sel yang di template memang kosong dibiarkan kosong: barisnya
            # berhenti di kolom tertentu, dan sel baru di kanannya akan muncul
            # tanpa border sebagai kolom liar.
            if ws[f"{kol}{nomor}"].value is None:
                continue
            lalu = _angka(total_lalu.get(NAMA_KOLOM[kol]))
            ws[f"{kol}{nomor}"] = lalu if kunci == "lalu" else nilai_ini - lalu

    buf = io.BytesIO()
    wb.save(buf)

    # Simpan setelah workbook selesai: laporan yang sudah jadi tidak boleh gagal
    # hanya karena volume arsipnya tidak bisa ditulis.
    try:
        simpan_arsip(bulan, tahun, total, per_kantor)
    except OSError as exc:
        log.warning("Arsip %s-%s gagal disimpan: %s", tahun, bulan, exc)
        catatan.append(f"Arsip bulan ini gagal disimpan ({str(exc)[:120]}) — "
                       f"laporan bulan depan tidak punya pembanding.")

    if kosong:
        catatan.append("Tanpa data di warehouse: " + ", ".join(sorted(kosong)))
    lebih = sorted(set(data) - terpakai)
    if lebih:
        # Baris kantor di template jumlahnya tetap. Kantor baru harus
        # ditambahkan barisnya di docs/template/summary.xlsx dulu, kalau tidak
        # angkanya hilang tanpa jejak.
        catatan.append("Ada di warehouse tapi tidak ada barisnya di template: "
                       + ", ".join(lebih))

    nama_file = f"rekap-{tahun:04d}-{bulan:02d}.xlsx"
    log.info("Laporan %s dibangun (%d kantor terisi, %d catatan).",
             nama_file, len(terpakai), len(catatan))
    return buf.getvalue(), nama_file, catatan
